"""Fill the Somer template copy (plantilla/base_23_diapositivas.pptx) with the talk's content.

Usage (from the repo root):  python presentacion/build_pptx.py   # needs: pip install python-pptx

Text (eyebrows, headlines, judge tags, footers, figure placeholders, speaker notes) is read
from the Slides-artifact HTML files so the .pptx and the web deck stay identical.
"""
import html
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SLIDES = HERE / "web" / "slides"
RENDER = HERE.parent / "figuras" / "render"          # S01.mp4 + S01.png, …, F03.png, F10.png, F11.png
VIDEO = RENDER / "S01.mp4"
POSTER = RENDER / "S01.png"
MEDIA = {"s03": "F03", "s04": "S04", "s06": "S06", "s07": "S07", "s09": "S09", "s11": "F11",
         "s16": "S16", "s17": "S17", "s18": "S18"}

ORDER = ["cover"] + [f"s{i:02d}" for i in range(1, 20)] + ["jueces", "s20", "s21"]

# ---- palette (template navy/blue + the talk's amber judge) ----
NAVY = RGBColor(0x07, 0x37, 0x63)
BLUE = RGBColor(0x15, 0x65, 0xC0)
BODY = RGBColor(0x4A, 0x4F, 0x57)
MUTED = RGBColor(0x75, 0x70, 0x70)
CARD = RGBColor(0xEE, 0xF3, 0xFA)
CARD_LINE = RGBColor(0xC9, 0xD8, 0xEE)
PH_FILL = RGBColor(0xF4, 0xF8, 0xFD)
PH_LINE = RGBColor(0x7F, 0xA7, 0xD9)
AMBER = RGBColor(0xF2, 0xA5, 0x41)
AMBER_TXT = RGBColor(0xA8, 0x60, 0x0A)
AMBER_TINT = RGBColor(0xFD, 0xF1, 0xDF)
WARN_TINT = RGBColor(0xFB, 0xEC, 0xE8)
WARN_TXT = RGBColor(0x8A, 0x3B, 0x26)
CODE_BG = RGBColor(0x0F, 0x16, 0x26)
CODE_TXT = RGBColor(0xDF, 0xE6, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
FONT = "Arial"
MONO = "Courier New"

X0, CW = 1.75, 16.5           # content left edge and width (in)
Y_EYE, Y_TITLE, Y_BODY = 1.85, 2.30, 4.25
Y_FOOT = 9.55


# ---------------------------------------------------------------- helpers
def parse(sid):
    h = (SLIDES / f"{sid}.html").read_text(encoding="utf-8")
    txt = lambda s: html.unescape(re.sub(r"<br\s*/?>", "\n", re.sub(r"<(?!br)[^>]+>", "", s))).strip()
    d = {"raw": h}
    m = re.search(r"<aside>(.*?)</aside>", h, re.S)
    d["notes"] = txt(m.group(1)) if m else ""
    m = re.search(r"<h[12][^>]*>(.*?)</h[12]>", h, re.S)
    d["title"] = txt(m.group(1)) if m else ""
    m = re.search(r"<section[^>]*>\s*<p[^>]*>(.*?)</p>", h, re.S)
    d["eyebrow"] = txt(m.group(1)) if m else ""
    m = re.search(r">Juez · <span[^>]*>(.*?)</span>", h, re.S)
    d["judge"] = txt(m.group(1)) if m else None
    m = re.search(r'<p style="position:absolute;left:128px;bottom:64px[^>]*>(.*?)</p>', h, re.S)
    d["footer"] = txt(m.group(1)) if m else ""
    m = re.search(r"dashed[^>]*>\s*<p[^>]*>(.*?)</p>\s*<p[^>]*>(.*?)</p>", h, re.S)
    d["ph"] = (txt(m.group(1)), txt(m.group(2))) if m else None
    return d


def clear(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def text(slide, x, y, w, h, paras, size=20, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing_after=0, line=None, name=None):
    """paras: str | list of (str | list of runs); a run is str or (str, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing_after:
            p.space_after = Pt(spacing_after)
        if line:
            p.line_spacing = line
        runs = para if isinstance(para, list) else [para]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = t
            f = run.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, radius=0.06, dash=False, name=None, lw=1.25):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            shp.line.dash_style = MSO_LINE.DASH
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def header(slide, d, title_size=40):
    if d["eyebrow"]:
        text(slide, X0, Y_EYE, 9.0, 0.4, d["eyebrow"].upper(), size=17, color=BLUE, bold=True, name="Eyebrow")
    if d["judge"]:
        text(slide, 7.6, Y_EYE, 6.4, 0.4,
             [[("JUEZ · ", {"color": MUTED}), (d["judge"].upper(), {"color": AMBER_TXT})]],
             size=17, bold=True, align=PP_ALIGN.RIGHT, name="Judge tag")
    text(slide, X0, Y_TITLE, CW, 1.75, d["title"], size=title_size, color=NAVY, bold=True,
         line=1.0, name="Title")


def footer(slide, d, color=MUTED, y=Y_FOOT, w=10.6):
    if d["footer"]:
        text(slide, X0, y, w, 0.45, d["footer"], size=16, color=color, name="Footer")


def placeholder(slide, d, x=X0, y=Y_BODY, w=CW, h=3.45, scene=None):
    tag, desc = d["ph"]
    is_scene = tag.startswith("ESCENA") if scene is None else scene
    box(slide, x, y, w, h, fill=PH_FILL, line=PH_LINE, dash=True, radius=0.04, name="Figure placeholder")
    pill_w = max(2.6, 0.17 * len(tag) + 0.6)
    pill = box(slide, x + (w - pill_w) / 2, y + h / 2 - 0.85, pill_w, 0.5,
               fill=AMBER if is_scene else BLUE, radius=0.5, name="Figure tag")
    text(slide, x + (w - pill_w) / 2, y + h / 2 - 0.85, pill_w, 0.5, tag, size=17, bold=True,
         color=NAVY if is_scene else WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.8, y + h / 2 - 0.15, w - 1.6, 1.2, desc, size=23, color=BODY,
         align=PP_ALIGN.CENTER, name="Figure description")


def media(slide, key, x, y, w, h):
    """Video (mp4 con su póster png) o figura png, en la caja dada."""
    mp4, png = RENDER / f"{key}.mp4", RENDER / f"{key}.png"
    if mp4.exists():
        return slide.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(h),
                                      poster_frame_image=str(png), mime_type="video/mp4")
    return slide.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(w), Inches(h))


def notes(slide, d):
    if d["notes"]:
        slide.notes_slide.notes_text_frame.text = d["notes"]


def card(slide, x, y, w, h, title, lines, title_color=NAVY, fill=CARD, line=None, body_size=22,
         title_size=26, sub=None):
    box(slide, x, y, w, h, fill=fill, line=line)
    paras = [[(title, {"size": title_size, "bold": True, "color": title_color})]]
    if sub:
        paras.append([(sub, {"size": 22, "bold": True, "color": AMBER_TXT})])
    for ln in lines:
        paras.append(ln if isinstance(ln, list) else [(ln, {})])
    text(slide, x + 0.35, y + 0.3, w - 0.7, h - 0.5, paras, size=body_size, color=BODY, spacing_after=8)


def bignum(slide, x, y, w, value, caption, color=NAVY, size=96, cap_size=20, source=None, rule=None):
    if rule is not None:
        ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = rule
        ln.line.width = Pt(3)
        y += 0.25
    vh = size / 72 * 1.05
    text(slide, x, y, w, vh, value, size=size, bold=True, color=color, line=0.9)
    text(slide, x, y + vh + 0.15, w, 1.4, caption, size=cap_size, color=BODY)
    if source:
        text(slide, x, y + vh + 1.45, w, 0.45, source, size=16, bold=True, color=MUTED)


# ---------------------------------------------------------------- slides
def build():
    prs = Presentation(HERE / "plantilla" / "base_23_diapositivas.pptx")
    D = {sid: parse(sid) for sid in ORDER}
    S = dict(zip(ORDER, prs.slides))
    for s in prs.slides:
        clear(s)

    # cover
    s, d = S["cover"], D["cover"]
    text(s, 1.71, 4.0, 12.0, 2.4,
         [[("¿Quién decide qué cuenta como ", {}), ("mejor", {"color": AMBER}), ("?", {})]],
         size=64, bold=True, color=WHITE, line=0.95, name="Title")
    text(s, 1.71, 6.45, 12.0, 0.9, "IA para la investigación: de la atención al descubrimiento autónomo",
         size=26, color=WHITE, name="Subtitle")
    text(s, 1.71, 7.45, 12.0, 0.5, "Juan David Martínez Vargas, Ph.D. · Universidad EAFIT",
         size=20, color=ICE, name="Author")
    notes(s, d)

    # s01 hook + video
    s, d = S["s01"], D["s01"]
    header(s, d)
    vw, vh = 11.3, 11.3 * 720 / 1920
    s.shapes.add_movie(str(VIDEO), Inches(X0), Inches(Y_BODY), Inches(vw), Inches(vh),
                       poster_frame_image=str(POSTER), mime_type="video/mp4")
    rx = X0 + vw + 0.5
    text(s, rx, Y_BODY, 18.25 - rx, 3.2, [
        [("El agente edita ", {}), ("train.py", {"font": MONO, "bold": True}), (", entrena 5 minutos y conserva el cambio solo si la validación mejora.", {})],
        [("Ámbar", {"bold": True, "color": AMBER_TXT}), (": lo que el juez conservó.", {})],
        [("Gris", {"bold": True, "color": MUTED}), (": lo que descartó.", {})],
    ], size=22, spacing_after=14)
    footer(s, d); notes(s, d)

    # s02 table of loops
    s, d = S["s02"], D["s02"]
    header(s, d)
    rows = re.findall(r"<tr>(.*?)</tr>", d["raw"], re.S)
    cells = [[html.unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, re.S)] for r in rows]
    tw = 15.6
    tbl = s.shapes.add_table(len(cells), 3, Inches(X0), Inches(4.1), Inches(tw), Inches(0.56 * len(cells))).table
    for j, frac in enumerate([0.31, 0.39, 0.30]):
        tbl.columns[j].width = Inches(tw * frac)
    for i, row in enumerate(cells):
        tbl.rows[i].height = Inches(0.56)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (CARD if i % 2 else WHITE)
            c.margin_left = c.margin_right = Inches(0.15)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = val
            r.font.name = FONT
            r.font.size = Pt(21)
            r.font.bold = i == 0 or (j == 2 and i == len(cells) - 1)
            r.font.color.rgb = WHITE if i == 0 else (AMBER_TXT if j == 2 and "mismos" not in val else BODY)
    m = re.search(r"</table>\s*<p[^>]*>(.*?)</p>", d["raw"], re.S)
    text(s, X0, 8.2, 12.6, 0.9, html.unescape(m.group(1)), size=23, bold=True, color=NAVY, name="Thesis")
    footer(s, d); notes(s, d)

    # escenas y figuras (8:3, oscuras, como una pantalla sobre la diapositiva blanca)
    for sid, key in MEDIA.items():
        s, d = S[sid], D[sid]
        header(s, d)
        mh = 4.75                      # alineado a la izquierda: deja libre la onda azul de la plantilla
        mw = mh * 1920 / 720
        media(s, key, X0, Y_BODY, mw, mh)
        footer(s, d); notes(s, d)

    # s05 three stages
    s, d = S["s05"], D["s05"]
    header(s, d)
    cw3 = (CW - 0.8) / 3
    stages = [
        ("1 · Preentrenamiento", "Juez: el siguiente token", "Billones de tokens de texto. El resultado completa documentos, pero no sigue instrucciones."),
        ("2 · Ajuste fino (SFT)", "Juez: ejemplos escritos por humanos", "Conversaciones modelo. Aprende el formato y el rol de asistente."),
        ("3 · RLHF", "Juez: las preferencias humanas", "¿Respuesta A o B? Moldea utilidad, honestidad y rechazos. No garantiza la verdad."),
    ]
    for k, (t, sub, body) in enumerate(stages):
        card(s, X0 + k * (cw3 + 0.4), Y_BODY, cw3, 3.4, t, [body], sub=sub)
    footer(s, d); notes(s, d)

    # s08 adoption big numbers
    s, d = S["s08"], D["s08"]
    header(s, d)
    bignum(s, X0, Y_BODY + 0.1, 6.8, "81%", "probó chatbots de IA para investigar", size=130, cap_size=26)
    bignum(s, X0 + 7.6, Y_BODY + 0.1, 6.8, "20%", "usa agentes de código más de una vez por semana",
           color=AMBER_TXT, size=130, cap_size=26)
    footer(s, d, w=14.5); notes(s, d)

    # s10 deep research inside: figure + three stats
    s, d = S["s10"], D["s10"]
    header(s, d)
    fw = 12.6
    media(s, "F10", X0, 4.15, fw, fw * 440 / 1920)
    stats = [("+90.2%", "mejor que un solo agente en su evaluación interna", AMBER_TXT),
             ("~15×", "más tokens que una conversación de chat", NAVY),
             ("80%", "del desempeño lo explica solo el uso de tokens", NAVY)]
    for k, (v, c, col) in enumerate(stats):
        x = X0 + k * 4.45
        text(s, x, 7.25, 4.2, 0.9, v, size=56, bold=True, color=col)
        text(s, x, 8.15, 4.2, 0.9, c, size=19, color=BODY)
    footer(s, d, w=11.5); notes(s, d)

    # s12 skill anatomy
    s, d = S["s12"], D["s12"]
    header(s, d)
    box(s, X0, Y_BODY, 8.0, 4.25, fill=CODE_BG, radius=0.04)
    code = [
        [("revisar-figura/", {})],
        [("├── ", {}), ("SKILL.md", {"color": AMBER, "bold": True})],
        [("├── checklist.md", {})],
        [("└── scripts/", {})],
        [("    └── contraste.py", {})],
        [("---", {})],
        [("name: ", {"color": AMBER}), ("revisar-figura", {})],
        [("description: ", {"color": AMBER}), ("Revisa una figura", {})],
        [("  científica (ejes, unidades, color,", {})],
        [("  leyenda) antes de enviarla.", {})],
        [("---", {})],
    ]
    text(s, X0 + 0.4, Y_BODY + 0.3, 7.3, 3.6, code, size=19, color=CODE_TXT, font=MONO, line=1.02)
    rx, rw = X0 + 8.45, CW - 8.45
    box(s, rx, Y_BODY, rw, 3.35, fill=CARD)
    levels = [[("Carga progresiva", {"size": 23, "bold": True, "color": NAVY})],
              [("1 · Al iniciar: ", {"bold": True, "color": AMBER_TXT}), ("solo el nombre y la descripción de cada skill", {})],
              [("2 · Si la tarea lo pide: ", {"bold": True, "color": AMBER_TXT}), ("el SKILL.md completo, con las instrucciones", {})],
              [("3 · Cuando hace falta: ", {"bold": True, "color": AMBER_TXT}), ("archivos de referencia y scripts, que corren sin llenar el contexto", {})]]
    text(s, rx + 0.35, Y_BODY + 0.3, rw - 0.7, 2.9, levels, size=21, spacing_after=10)
    footer(s, d, w=11.5); notes(s, d)

    # s13 five steps
    s, d = S["s13"], D["s13"]
    header(s, d)
    steps = ["Elige algo que ya hiciste tres veces", "Escribe los pasos y los errores típicos",
             "Describe cuándo usarlo: eso lo activa", "Pasa lo mecánico a un script",
             "Pruébalo en tres casos reales e itera"]
    sw = (CW - 4 * 0.3) / 5
    for k, st in enumerate(steps):
        x = X0 + k * (sw + 0.3)
        box(s, x, Y_BODY, sw, 2.6, fill=CARD)
        text(s, x + 0.3, Y_BODY + 0.2, sw - 0.6, 0.8, str(k + 1), size=44, bold=True, color=AMBER_TXT)
        text(s, x + 0.3, Y_BODY + 1.0, sw - 0.6, 1.4, st, size=22, color=NAVY)
    box(s, X0, 7.15, 13.9, 0.8, fill=CODE_BG, radius=0.12)
    text(s, X0 + 0.35, 7.15, 13.3, 0.8, "Demo · «Crea un skill revisar-figura con estos pasos y pruébalo con esta figura»",
         size=18, color=CODE_TXT, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, d, w=12.5); notes(s, d)

    # s14 gallery
    s, d = S["s14"], D["s14"]
    header(s, d)
    gallery = [("Revisar papers", "/review-paper de Bäckman: varios agentes revisores escriben un reporte de árbitro en 10 a 15 minutos (licencia MIT)"),
               ("Ciencias de la vida", "claude-scientific-skills de K-Dense: conecta al agente con bases como PubMed, ChEMBL y UniProt"),
               ("Experimentos", "autoresearch empaquetado como skill: el ciclo de proponer, correr y evaluar para cualquier métrica")]
    for k, (t, b) in enumerate(gallery):
        card(s, X0 + k * (cw3 + 0.4), Y_BODY, cw3, 3.05, t, [b], title_color=BLUE)
    box(s, X0, 7.55, 12.8, 0.85, fill=WARN_TINT, radius=0.12)
    text(s, X0 + 0.35, 7.55, 12.2, 0.85, "Instala solo de fuentes confiables y léelos antes: un skill puede ordenarle al agente acciones que no quieres.",
         size=20, color=WARN_TXT, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, d, w=13.5); notes(s, d)

    # s15 harness = judge
    s, d = S["s15"], D["s15"]
    header(s, d)
    practices = [("Un mapa, no un manual", "Un CLAUDE.md o AGENTS.md corto que apunta a lo demás"),
                 ("Restricciones mecánicas", "Tests y linters que el agente no puede saltarse"),
                 ("Evaluar el comportamiento", "¿Corrió el validador? ¿Preguntó antes de asumir?"),
                 ("Un revisor independiente", "Otro agente, u otro modelo, ataca el resultado")]
    hw = (CW - 0.4) / 2
    for k, (t, b) in enumerate(practices):
        card(s, X0 + (k % 2) * (hw + 0.4), Y_BODY + (k // 2) * 1.9, hw, 1.7, t, [b], title_size=24, body_size=21)
    text(s, X0, 8.15, 13.2, 0.9, "En autoresearch, el harness es el experimento: program.md + train.py + una métrica.",
         size=24, bold=True, color=AMBER_TXT)
    footer(s, d, w=12.5); notes(s, d)

    # s19 three end-to-end numbers
    s, d = S["s19"], D["s19"]
    header(s, d)
    cols = [("6.33", "puntaje del primer paper 100% IA que pasó revisión en un workshop de ICLR 2025", "The AI Scientist · Nature 2026", NAVY, CARD_LINE),
            ("79.4%", "de las afirmaciones de Kosmos son correctas: cerca de 1 de cada 5 no lo es", "Kosmos · arXiv 2511.02824", NAVY, CARD_LINE),
            ("88.1%", "de las veces, un juez prefiere el paper humano frente a los ~1,000 papers de Project APE", "Project APE · Universidad de Zurich", AMBER_TXT, AMBER)]
    for k, (v, c, src, col, rule) in enumerate(cols):
        bignum(s, X0 + k * 5.2, Y_BODY, 4.7, v, c, color=col, size=80, cap_size=21, source=src, rule=rule)
    footer(s, d); notes(s, d)

    # jueces: judge of the judge
    s, d = S["jueces"], D["jueces"]
    header(s, d)
    text(s, X0, 3.95, CW, 0.5, "¿Con qué frecuencia cada juez coincide con la respuesta que prefieren los expertos?",
         size=22, color=BODY, bold=True)
    jcols = [("58%", "Una métrica de similitud (ROUGE). El azar da 50%.", MUTED, CARD_LINE),
             ("74%", "Un LLM como juez, guiado por rúbricas.", AMBER_TXT, AMBER),
             ("84%", "Dos expertos entre sí: el techo realista.", NAVY, NAVY)]
    for k, (v, c, col, rule) in enumerate(jcols):
        bignum(s, X0 + k * 5.2, 4.65, 4.7, v, c, color=col, size=88, cap_size=22, rule=rule)
    box(s, X0, 7.85, 12.9, 1.05, fill=WARN_TINT, radius=0.1)
    text(s, X0 + 0.35, 7.85, 12.3, 1.05,
         "Y la rúbrica también falla: el 11% de los criterios generados por un modelo, sin fuentes, eran vacíos, erróneos o poco claros. Audita al juez antes de confiar en él.",
         size=19, color=WARN_TXT, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, d, w=13.0); notes(s, d)

    # s20 failures + checks
    s, d = S["s20"], D["s20"]
    header(s, d)
    card(s, X0, Y_BODY, hw, 3.45, "Cómo se rompe",
         [[("Sin juez: ", {"bold": True, "color": NAVY}), ("citas inventadas", {})],
          [("Juez engañado: ", {"bold": True, "color": NAVY}), ("reward hacking", {})],
          [("Juez con fuga: ", {"bold": True, "color": NAVY}), ("datos de prueba en el entrenamiento", {})],
          [("Entrada envenenada: ", {"bold": True, "color": NAVY}), ("prompt injection", {})]])
    card(s, X0 + hw + 0.4, Y_BODY, hw, 3.2, "Cuatro checks que ayudan",
         ["Un agente nuevo ataca los cambios", "Un modelo de otro laboratorio lo revisa",
          "El agente te hace preguntas sobre el cambio", "Tú reimplementas el resultado principal"],
         title_color=AMBER_TXT, fill=AMBER_TINT, line=AMBER)
    footer(s, d, w=13.0); notes(s, d)

    # s21 closing on the cover background
    s, d = S["s21"], D["s21"]
    text(s, 1.71, 3.05, 10.0, 0.4, d["eyebrow"].upper(), size=15, bold=True, color=ICE)
    text(s, 1.71, 3.6, 12.2, 2.3, d["title"], size=52, bold=True, color=WHITE, line=0.95)
    tiers = [("1 · Contexto", "Un CLAUDE.md con tu campo y tus reglas, y tus propios papers como contexto"),
             ("2 · Skills", "Empaqueta lo que repites; empieza por /review-paper sobre tu próximo borrador"),
             ("3 · Ciclos", "Ciclos autónomos solo donde tengas una métrica barata y honesta")]
    tw3 = 4.9
    for k, (t, b) in enumerate(tiers):
        x = 1.71 + k * (tw3 + 0.4)
        ln = s.shapes.add_connector(1, Inches(x), Inches(5.85), Inches(x + tw3), Inches(5.85))
        ln.line.color.rgb = AMBER
        ln.line.width = Pt(3)
        text(s, x, 6.05, tw3, 1.5, [[(t, {"size": 25, "bold": True, "color": AMBER})], [(b, {})]],
             size=20, color=WHITE, spacing_after=6)
    text(s, 1.71, 7.88, 12.0, 0.4, d["footer"], size=16, color=ICE)
    notes(s, d)

    out = HERE / "IA-para-la-investigacion_Simposio-Somer.pptx"
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    build()
