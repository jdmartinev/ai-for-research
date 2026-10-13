"""spec.py → .pptx sobre la plantilla del 1er Simposio Somer Investiga.

python render_pptx.py [plantilla.pptx] [carpeta_de_renders] [salida.pptx]
Necesita python-pptx y ffprobe (para la duración de los videos, que arrancan solos).
"""
import copy
import re
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import spec  # noqa: E402

BASE = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "plantilla" / "base_23_diapositivas.pptx"
RENDER = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE.parent / "figuras" / "render"
OUT = Path(sys.argv[3]) if len(sys.argv) > 3 else HERE / "IA-para-la-investigacion_Simposio-Somer.pptx"

NAVY, BLUE, BODY, MUTED = RGBColor(0x07, 0x37, 0x63), RGBColor(0x15, 0x65, 0xC0), RGBColor(0x4A, 0x4F, 0x57), RGBColor(0x75, 0x70, 0x70)
CARD, CARD_LINE, PH_FILL, PH_LINE = RGBColor(0xEE, 0xF3, 0xFA), RGBColor(0xC9, 0xD8, 0xEE), RGBColor(0xF4, 0xF8, 0xFD), RGBColor(0x7F, 0xA7, 0xD9)
AMBER, AMBER_TXT, AMBER_TINT = RGBColor(0xF2, 0xA5, 0x41), RGBColor(0xA8, 0x60, 0x0A), RGBColor(0xFD, 0xF1, 0xDF)
WARN_TINT, WARN_TXT = RGBColor(0xFB, 0xEC, 0xE8), RGBColor(0x8A, 0x3B, 0x26)
CODE_BG, CODE_TXT, CODE_GRAY, CODE_BLUE = RGBColor(0x0F, 0x16, 0x26), RGBColor(0xDF, 0xE6, 0xEE), RGBColor(0x8A, 0x96, 0xA8), RGBColor(0x5F, 0xA8, 0xD3)
WHITE, ICE = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xCA, 0xDC, 0xFC)
ROLE = {"a": AMBER, "m": CODE_BLUE, "g": CODE_GRAY, "t": CODE_TXT}
FONT, MONO = "Arial", "Courier New"
X0, CW, Y_EYE, Y_TITLE, Y_BODY, Y_FOOT = 1.75, 16.5, 1.85, 2.30, 4.25, 9.55
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

ORDER = [d["id"] for d in spec.SLIDES]
NUM = {sid: i for i, sid in enumerate(ORDER)}


def refs(s):
    return re.sub(r"\{n:([\w-]+)\}", lambda m: str(NUM[m.group(1)]), s)


# ---------------------------------------------------------------- plantilla
def build_deck():
    """Copia la portada (diap. 1), una de contenido (diap. 2) y la de cierre (última) tantas veces como haga falta."""
    prs = Presentation(BASE)
    slides = list(prs.slides)
    proto = {"cover": slides[0], "content": slides[1], "closing": slides[-1]}
    made = []
    for d in spec.SLIDES:
        kind = "cover" if d["kind"] == "cover" else ("closing" if d["kind"] == "closing" else "content")
        made.append(dup(prs, proto[kind]))
    # quitar las 23 originales
    lst = prs.slides._sldIdLst
    for sld in slides:
        for sid in list(lst):
            if prs.part.related_part(sid.get(f"{{{R_NS}}}id")) is sld.part:
                prs.part.drop_rel(sid.get(f"{{{R_NS}}}id"))
                lst.remove(sid)
    return prs, made


def dup(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.bg
    if bg is not None:
        nb = copy.deepcopy(bg)
        for blip in nb.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            old = blip.get(f"{{{R_NS}}}embed")
            part = src.part.related_part(old)
            blip.set(f"{{{R_NS}}}embed", new.part.relate_to(part, src.part.rels[old].reltype))
        new._element.cSld.insert(0, nb)
    return new


# ---------------------------------------------------------------- piezas
def text(slide, x, y, w, h, paras, size=20, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         font=FONT, spacing_after=0, line=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing_after:
            p.space_after = Pt(spacing_after)
        if line:
            p.line_spacing = line
        for r in (para if isinstance(para, list) else [para]):
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = refs(t)
            f = run.font
            f.name, f.size, f.bold = o.get("font", font), Pt(o.get("size", size)), o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, radius=0.06, dash=False, lw=1.25):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def rule(slide, x, y, w, color, width=3):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)


def header(s, d):
    text(s, X0, Y_EYE, 9.0, 0.4, d["eyebrow"].upper(), size=17, color=BLUE, bold=True)
    if d.get("judge"):
        text(s, 7.6, Y_EYE, 6.4, 0.4, [[("JUEZ · ", {"color": MUTED}), (d["judge"].upper(), {"color": AMBER_TXT})]],
             size=17, bold=True, align=PP_ALIGN.RIGHT)
    text(s, X0, Y_TITLE, CW, 1.75, d["title"], size=38, color=NAVY, bold=True, line=1.0)


def footer(s, d, w=12.0):
    text(s, X0, Y_FOOT, w, 0.45, f"{NUM[d['id']]} · {d.get('src', 'IA para la investigación')}", size=15, color=MUTED)


def notes(s, d):
    if d.get("notes"):
        s.notes_slide.notes_text_frame.text = refs(d["notes"])


AUTOPLAY = """<p:timing xmlns:p="{ns}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{dur}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def media(slide, key, x, y, w, h, alt=""):
    mp4, png = RENDER / f"{key}.mp4", RENDER / f"{key}.png"
    if mp4.exists():
        mv = slide.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(h),
                                    poster_frame_image=str(png), mime_type="video/mp4")
        dur = float(subprocess.check_output(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                                             "-of", "csv=p=0", str(mp4)]).decode())
        sld = slide._element
        for old in sld.findall(f"{{{P_NS}}}timing"):
            sld.remove(old)
        timing = etree.fromstring(AUTOPLAY.format(ns=P_NS, dur=int(dur * 1000), spid=mv.shape_id))
        ext = sld.find(f"{{{P_NS}}}extLst")
        (ext.addprevious(timing) if ext is not None else sld.append(timing))
        return mv
    if png.exists():
        return slide.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(w), Inches(h))
    box(slide, x, y, w, h, fill=PH_FILL, line=PH_LINE, dash=True, radius=0.04)
    text(slide, x + 0.6, y + h / 2 - 0.6, w - 1.2, 1.2, [[(f"EN CONSTRUCCIÓN · {key}", {"bold": True, "color": AMBER_TXT, "size": 18})], alt],
         size=20, align=PP_ALIGN.CENTER)


def card(s, x, y, w, h, title, body, sub=None, accent=False, title_size=24, body_size=20):
    box(s, x, y, w, h, fill=AMBER_TINT if accent else CARD, line=AMBER if accent else None)
    paras = [[(title, {"size": title_size, "bold": True, "color": AMBER_TXT if accent else NAVY})]]
    if sub:
        paras.append([(sub, {"size": body_size, "bold": True, "color": AMBER_TXT})])
    for b in (body if isinstance(body, list) else [body]):
        paras.append(b if isinstance(b, list) else [(b, {})])
    text(s, x + 0.3, y + 0.25, w - 0.6, h - 0.4, paras, size=body_size, spacing_after=8)


def note_warn(s, d, y):
    if d.get("note"):
        text(s, X0, y, 13.0, 0.9, d["note"], size=22, bold=True, color=AMBER_TXT)
        y += 0.9
    if d.get("warn"):
        box(s, X0, y, 13.0, 0.95, fill=WARN_TINT, radius=0.1)
        text(s, X0 + 0.3, y, 12.4, 0.95, d["warn"], size=18, color=WARN_TXT, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- kinds
def k_cover(s, d):
    text(s, 1.71, 4.0, 12.0, 2.4, [[("¿Quién decide qué cuenta como ", {}), ("mejor", {"color": AMBER}), ("?", {})]],
         size=64, bold=True, color=WHITE, line=0.95)
    text(s, 1.71, 6.45, 12.0, 0.9, "IA para la investigación: de la atención al descubrimiento autónomo", size=26, color=WHITE)
    text(s, 1.71, 7.45, 12.0, 0.5, "Juan David Martínez Vargas, Ph.D. · Universidad EAFIT", size=20, color=ICE)


def k_media(s, d):
    header(s, d)
    mh = 4.75
    media(s, d["media"], X0, Y_BODY, mh * 1920 / 720, mh, d.get("alt", ""))
    footer(s, d)


def k_media_stats(s, d):
    header(s, d)
    fw = 12.6
    media(s, d["media"], X0, 4.15, fw, fw * 440 / 1920, d.get("alt", ""))
    for k, (v, c, acc) in enumerate(d["stats"]):
        x = X0 + k * 4.45
        text(s, x, 7.25, 4.2, 0.9, v, size=56, bold=True, color=AMBER_TXT if acc else NAVY)
        text(s, x, 8.15, 4.2, 0.9, c, size=19)
    footer(s, d, w=11.5)


def k_table(s, d):
    header(s, d)
    rows = [d["head"]] + d["rows"]
    tw = 15.6
    tbl = s.shapes.add_table(len(rows), 3, Inches(X0), Inches(4.1), Inches(tw), Inches(0.56 * len(rows))).table
    for j, frac in enumerate([0.31, 0.39, 0.30]):
        tbl.columns[j].width = Inches(tw * frac)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(0.56)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (CARD if i % 2 else WHITE)
            c.margin_left = c.margin_right = Inches(0.15)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = val
            r.font.name, r.font.size = FONT, Pt(21)
            r.font.bold = i == 0 or (j == 2 and i == len(rows) - 1)
            r.font.color.rgb = WHITE if i == 0 else (AMBER_TXT if j == 2 else BODY)
    text(s, X0, 8.2, 12.6, 0.9, d["note"], size=23, bold=True, color=NAVY)
    footer(s, d)


def k_cards(s, d):
    header(s, d)
    cs = d["cards"]
    cols = d.get("grid", len(cs))
    rows = (len(cs) + cols - 1) // cols
    gap = 0.35
    w = (CW - gap * (cols - 1)) / cols
    h = 3.2 if rows == 1 else 1.55
    for k, c in enumerate(cs):
        x = X0 + (k % cols) * (w + gap)
        y = Y_BODY + (k // cols) * (h + 0.25)
        card(s, x, y, w, h, c["t"], c["body"], sub=c.get("sub"), accent=c.get("accent"),
             title_size=24 if rows == 1 else 22, body_size=20 if rows == 1 else 18)
    note_warn(s, d, Y_BODY + rows * (h + 0.25) + 0.15)
    footer(s, d)


def k_bignums(s, d):
    header(s, d)
    y = Y_BODY
    if d.get("lead"):
        text(s, X0, 3.95, CW, 0.5, d["lead"], size=22, bold=True)
        y = 4.65
    items = d["items"]
    n = len(items)
    big = any(i.get("big") for i in items)
    w = (CW - 0.6 * (n - 1)) / n
    L = max(len(i["v"]) for i in items)
    size = 130 if big else (80 if L <= 6 else (54 if L <= 10 else 40))
    for k, it in enumerate(items):
        x = X0 + k * (w + 0.6)
        col = AMBER_TXT if it.get("accent") else (MUTED if it.get("muted") else NAVY)
        rule(s, x, y, w, AMBER if it.get("accent") else CARD_LINE)
        vh = size / 72 * 1.05
        text(s, x, y + 0.25, w, vh, it["v"], size=size, bold=True, color=col, line=0.9)
        text(s, x, y + 0.25 + vh + 0.1, w, 1.5, it["cap"], size=24 if big else 20)
        if it.get("src"):
            text(s, x, y + 0.25 + vh + 1.45, w, 0.45, it["src"], size=15, bold=True, color=MUTED)
    note_warn(s, d, 7.9)
    footer(s, d, w=13.0)


def k_compare(s, d):
    header(s, d)
    w = (CW - 0.4) / 2
    for k, side in enumerate((d["left"], d["right"])):
        card(s, X0 + k * (w + 0.4), Y_BODY, w, 3.9, side["t"], [[("•  " + it, {})] for it in side["items"]],
             accent=side.get("accent"), body_size=21)
    note_warn(s, d, 8.3)
    footer(s, d, w=13.0)


def k_steps(s, d):
    header(s, d)
    n = len(d["steps"])
    w = (CW - 0.3 * (n - 1)) / n
    for k, st in enumerate(d["steps"]):
        x = X0 + k * (w + 0.3)
        box(s, x, Y_BODY, w, 2.6, fill=CARD)
        text(s, x + 0.3, Y_BODY + 0.2, w - 0.6, 0.8, str(k + 1), size=44, bold=True, color=AMBER_TXT)
        text(s, x + 0.3, Y_BODY + 1.0, w - 0.6, 1.5, st, size=21 if n > 4 else 22, color=NAVY)
    y = Y_BODY + 2.9
    if d.get("demo"):
        box(s, X0, y, 13.9, 0.8, fill=CODE_BG, radius=0.12)
        text(s, X0 + 0.35, y, 13.3, 0.8, d["demo"], size=18, color=CODE_TXT, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.0
    note_warn(s, d, y)
    footer(s, d, w=13.0)


def k_code(s, d):
    header(s, d)
    cw = 10.2
    n = len(d["code"])
    box(s, X0, Y_BODY, cw, 4.8, fill=CODE_BG, radius=0.04)
    paras = []
    for ln in d["code"]:
        if isinstance(ln, str):
            paras.append([(ln or " ", {})])
        else:
            paras.append([(t, {"color": ROLE[r]}) for t, r in ln])
    text(s, X0 + 0.35, Y_BODY + 0.3, cw - 0.6, 4.3, paras, size=17 if n <= 10 else 15, color=CODE_TXT, font=MONO, line=1.05)
    rx, rw = X0 + cw + 0.35, CW - cw - 0.35
    box(s, rx, Y_BODY, rw, 4.8, fill=CARD)
    side = []
    for a, b in d["side"]:
        side.append([(a, {"bold": True, "color": AMBER_TXT, "size": 21})])
        side.append([(b, {})])
    text(s, rx + 0.3, Y_BODY + 0.3, rw - 0.6, 4.3, side, size=18, spacing_after=6)
    footer(s, d, w=13.0)


def k_timeline(s, d):
    header(s, d)
    n = len(d["items"])
    w = (CW - 0.4 * (n - 1)) / n
    for k, (when, what) in enumerate(d["items"]):
        x = X0 + k * (w + 0.4)
        last = k == n - 1
        rule(s, x, Y_BODY + 0.2, w, AMBER if last else CARD_LINE)
        text(s, x, Y_BODY + 0.45, w, 0.9, when, size=40, bold=True, color=AMBER_TXT if last else NAVY)
        text(s, x, Y_BODY + 1.45, w, 2.4, what, size=21)
    footer(s, d, w=13.5)


def k_closing(s, d):
    text(s, 1.71, 3.05, 10.0, 0.4, d["eyebrow"].upper(), size=15, bold=True, color=ICE)
    text(s, 1.71, 3.6, 12.2, 2.3, d["title"], size=52, bold=True, color=WHITE, line=0.95)
    tw = 4.9
    for k, c in enumerate(d["cards"]):
        x = 1.71 + k * (tw + 0.4)
        rule(s, x, 5.85, tw, AMBER)
        text(s, x, 6.05, tw, 1.5, [[(c["t"], {"size": 25, "bold": True, "color": AMBER})], [(c["body"], {})]],
             size=20, color=WHITE, spacing_after=6)
    text(s, 1.71, 7.88, 12.0, 0.4, f"{NUM[d['id']]} · {d['src']}", size=16, color=ICE)


KINDS = {k[2:]: v for k, v in globals().items() if k.startswith("k_")}


def main():
    prs, made = build_deck()
    for s, d in zip(made, spec.SLIDES):
        KINDS[d["kind"]](s, d)
        notes(s, d)
    prs.save(OUT)
    print("saved", OUT, len(made), "slides")


if __name__ == "__main__":
    main()
